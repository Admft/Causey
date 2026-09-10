#!/usr/bin/env python3
"""Internal partnerships + marketing briefing. Not a customer deck."""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_template import (  # noqa: E402
    ACCENT_SOFT,
    ASSETS,
    BG,
    BLUE,
    BLUE_STRONG,
    FONT_DISPLAY,
    FONT_SANS,
    FORE,
    LINE,
    MUTED,
    RED,
    ROOT,
    SLIDE_H,
    SLIDE_W,
    SOFT,
    SURFACE,
    _set_solid,
    add_master_chrome,
    add_run,
    add_textbox,
    apply_slide_footer_fields,
    card,
    patch_master_text_styles,
    patch_theme,
    set_box_text,
    style_layouts,
    write_assets,
)

SHOTS = ROOT / "shots"
OUT = ROOT / "Causey team briefing.pptx"


def notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def picture(slide, name: str, left, top, width):
    return slide.shapes.add_picture(str(SHOTS / name), left, top, width=width)


def caption(slide, left, top, width, text: str):
    add_textbox(
        slide, left, top, width, Inches(0.32),
        [text], font=FONT_SANS, size=11, color=MUTED,
    )


def new_slide(prs, layout_idx=6):
    slide = prs.slides.add_slide(prs.slide_layouts[layout_idx])
    apply_slide_footer_fields(slide)
    return slide


def title_bar(slide, title: str):
    set_box_text(
        slide.placeholders[0], [title],
        font=FONT_DISPLAY, size=28, color=FORE, bold=True,
    )


def build(prs, assets) -> None:
    # 1. Title
    slide = new_slide(prs, 0)
    slide.shapes.add_picture(str(assets["lockup"]), Inches(0.72), Inches(1.2), height=Inches(0.5))
    set_box_text(
        slide.placeholders[0],
        ["What’s built: a short product overview"],
        font=FONT_DISPLAY, size=36, color=FORE, bold=True,
    )
    set_box_text(
        slide.placeholders[1],
        ["CTO brief  ·  partnerships and marketing  ·  7 September 2026"],
        font=FONT_SANS, size=16, color=MUTED,
    )
    notes(slide, "You are describing the product as it exists. GTM choices sit with the founder.")

    # 2. Why this overview exists
    slide = new_slide(prs, 5)
    title_bar(slide, "How this maps to your work")
    card(
        slide, Inches(0.72), Inches(1.5), Inches(5.9), Inches(5.0),
        "Partnerships",
        "District packages sit on the assisted chess workspace: we provision, schools run events, the office sees totals. Competition packages sit on the public index: chess is already listed from permitted hubs; other categories stay thin until those organizations allow us to index public dates.",
    )
    card(
        slide, Inches(6.85), Inches(1.5), Inches(5.75), Inches(5.0),
        "Social / marketing",
        "People can search without an account. Accounts unlock save, RSVP, family follow-through, and club or school coordination. That is the surface area of the product today.",
        bar_color=BLUE, label_color=BLUE_STRONG,
    )
    notes(slide, "This is context, not an assignment list.")

    # 3. What Causey is + screenshot
    slide = new_slide(prs, 5)
    title_bar(slide, "The website, today")
    picture(slide, "web-home.png", Inches(0.72), Inches(1.42), Inches(7.55))
    caption(slide, Inches(0.72), Inches(6.72), Inches(7.55), "causey.dev  ·  guest search, no account required")
    add_textbox(
        slide, Inches(8.5), Inches(1.5), Inches(4.15), Inches(5.2),
        [
            "Student competitions, indexed in one place.",
            "",
            "Search by type, zip, and distance with no account.",
            "",
            "Chess is the densest directory. Debate, STEM, arts, and writing have search pages and are still thin.",
            "",
            "Paid entry stays on the organizer’s site. Causey does not collect registration fees.",
        ],
        font=FONT_SANS, size=15, color=FORE,
    )
    notes(slide, "Coverage is incomplete even for chess.")

    # 4. Accounts for marketing
    slide = new_slide(prs, 5)
    title_bar(slide, "What an account unlocks")
    add_textbox(
        slide, Inches(0.72), Inches(1.35), Inches(12.0), Inches(0.4),
        ["Useful if you are describing signup on social or in a district conversation."],
        font=FONT_SANS, size=14, color=MUTED,
    )
    rows = [
        ("No account", "Search public listings. Chess is usable. Other types may return few or no events."),
        ("Student", "Save events, RSVP, keep a plan, join a club or school with a code. Student accounts are created on the website."),
        ("Parent", "Link a child. Family desk: invite, going / can’t go, mark organizer registration done, alerts."),
        ("Coach / club owner", "Create a club or team. Roster, travel (“we are going”), host an event, attendance, results, season CSV."),
        ("District / school", "Causey sets the district up. Not a public signup. School staff invite their own people. District office sees school totals, not browsing history."),
    ]
    y = Inches(1.78)
    for label, body in rows:
        bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.72), y, Inches(11.9), Inches(0.88))
        _set_solid(bar, SURFACE, LINE)
        try:
            bar.adjustments[0] = 0.12
        except Exception:
            pass
        add_textbox(slide, Inches(0.92), y + Inches(0.12), Inches(2.5), Inches(0.64), [label], font=FONT_SANS, size=14, color=RED, bold=True)
        add_textbox(slide, Inches(3.5), y + Inches(0.12), Inches(8.85), Inches(0.64), [body], font=FONT_SANS, size=14, color=FORE)
        y += Inches(0.98)
    notes(slide, "Phone signup is parent or coach. Student accounts, including under 13, are a website path.")

    # 5. Club screenshot
    slide = new_slide(prs, 5)
    title_bar(slide, "Clubs and teams")
    picture(slide, "web-clubs.png", Inches(0.72), Inches(1.42), Inches(7.55))
    caption(slide, Inches(0.72), Inches(6.72), Inches(7.55), "causey.dev/clubs  ·  self-serve club or team")
    add_textbox(
        slide, Inches(8.5), Inches(1.5), Inches(4.15), Inches(5.2),
        [
            "A club or team workspace, created by the coach.",
            "",
            "Season path: roster → find or host → invite / RSVP → attendance and place → season CSV.",
            "",
            "Chess clubs see the most public listings. Other types can still host and run a roster.",
            "",
            "Not in the product: pairings, student dues, or a public club directory.",
        ],
        font=FONT_SANS, size=15, color=FORE,
    )
    notes(slide, "Club SaaS checkout is a local layout. Nothing is charging yet.")

    # 6. District screenshot
    slide = new_slide(prs, 5)
    title_bar(slide, "Schools and districts")
    picture(slide, "web-districts.png", Inches(0.72), Inches(1.42), Inches(7.55))
    caption(slide, Inches(0.72), Inches(6.72), Inches(7.55), "causey.dev/districts  ·  assisted chess pilot")
    add_textbox(
        slide, Inches(8.5), Inches(1.5), Inches(4.15), Inches(5.2),
        [
            "Assisted chess pilot. Causey creates the district; named school admins claim; coaches run events; the office reads school-level totals, not browsing history.",
            "",
            "Other types can be hosted in a school workspace. Chess is the working surface.",
            "",
            "There is no self-serve district signup and no public school directory. Price, SLA, and privacy certification are not finished product.",
        ],
        font=FONT_SANS, size=15, color=FORE,
    )
    notes(slide, "Public search is free. The district workspace is coordination on top of that.")

    # 7. Phone + stores
    slide = new_slide(prs, 5)
    title_bar(slide, "iOS app submitted  ·  Android next")
    picture(slide, "phone-home.png", Inches(0.72), Inches(1.4), Inches(2.55))
    picture(slide, "phone-chess.png", Inches(3.4), Inches(1.4), Inches(2.55))
    caption(slide, Inches(0.72), Inches(6.72), Inches(5.2), "Website on a phone. Native iOS is in review.")
    add_textbox(
        slide, Inches(6.3), Inches(1.5), Inches(6.3), Inches(5.2),
        [
            "iOS is submitted and in review. Android is not in a store yet.",
            "",
            "Native app, same account as the website: search, save, family RSVP, coach team / attendance / results, alerts, join code.",
            "",
            "Still website-only: club setup, CSV invites, district provisioning, season reports, student signup.",
            "",
            "The app does not run registration. It tracks that the family finished it on the organizer’s site.",
        ],
        font=FONT_SANS, size=15, color=FORE,
    )
    notes(slide, "Phone screenshots here are the website at phone width. Native iOS is the App Store binary.")

    # 8. Biggest features
    slide = new_slide(prs, 5)
    title_bar(slide, "What is working in the product")
    items = [
        ("Public chess search", "Zip, distance, date, and related filters. Sources include US Chess TLA and other permitted chess hubs. Still incomplete."),
        ("Family follow-through", "Invite → plan → going / can’t go → mark organizer registration complete. Alerts on web and phone."),
        ("Club or team season", "Roster, travel and hosted events, attendance, recorded place/award, season CSV."),
        ("District command center", "Schools, claim links, hosted school and district events, RSVP by school, aggregate reports. Causey-assisted."),
        ("Honest empty states", "Thin directories say so. Listings, fees, and partner names are not invented in the software."),
    ]
    y = Inches(1.48)
    for i, (h, b) in enumerate(items, 1):
        add_textbox(slide, Inches(0.72), y, Inches(0.45), Inches(0.9), [f"0{i}"], font=FONT_DISPLAY, size=20, color=RED, bold=True)
        add_textbox(slide, Inches(1.28), y, Inches(11.2), Inches(0.38), [h], font=FONT_SANS, size=16, color=FORE, bold=True)
        add_textbox(slide, Inches(1.28), y + Inches(0.36), Inches(11.2), Inches(0.5), [b], font=FONT_SANS, size=14, color=MUTED)
        y += Inches(1.0)
    notes(slide, "Chess pathways exist as illustrative scaffolding, not an official US Chess ruling.")

    # 9. Coverage permissions
    slide = new_slide(prs, 5)
    title_bar(slide, "Where the index stops today")
    add_textbox(
        slide, Inches(0.72), Inches(1.38), Inches(11.9), Inches(0.7),
        ["Chess is the working public index. Other categories stay thin until these organizations allow us to republish public listing facts. That is an engineering constraint, not a partnership announcement."],
        font=FONT_SANS, size=15, color=FORE,
    )
    table_shape = slide.shapes.add_table(5, 3, Inches(0.72), Inches(2.2), Inches(11.9), Inches(4.35))
    table = table_shape.table
    table.columns[0].width = Inches(2.8)
    table.columns[1].width = Inches(3.5)
    table.columns[2].width = Inches(5.6)
    data = [
        ["Org", "What it would add to search", "What engineering needs to ingest"],
        ["NSDA / Tabroom", "Nearly every debate tournament", "License to fetch public listings, including automated access and commercial reuse"],
        ["Scholastic Art & Writing", "National arts and writing program", "Index public program and affiliate deadlines, or a feed they already publish"],
        ["MATHCOUNTS", "Middle-school math, chapter through nationals", "Written consent to republish public competition search results"],
        ["FIRST (FRC / FTC / FLL)", "Robotics at school age bands", "Listings license for events only, with attribution. Their free API token does not cover a paid product"],
    ]
    for r, row in enumerate(data):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = val
            cell.margin_left = Inches(0.1)
            cell.margin_top = Inches(0.08)
            fill = cell.fill
            fill.solid()
            fill.fore_color.rgb = RED if r == 0 else (SURFACE if r % 2 else SOFT)
            color = SURFACE if r == 0 else FORE
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT
                for run in p.runs:
                    run.font.name = FONT_SANS
                    run.font.size = Pt(12)
                    run.font.bold = r == 0 or c == 0
                    run.font.color.rgb = color
    notes(slide, "Next ring if those land: SpeechWire, MAA, VEX/REC, Society for Science. Science Bowl is already indexed from DOE public pages. Adapters stay off until written permission is on file.")

    # 10. Monetization later
    slide = new_slide(prs, 5)
    title_bar(slide, "What is next on monetization")
    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.72), Inches(1.5), Inches(11.9), Inches(1.15))
    _set_solid(panel, ACCENT_SOFT)
    try:
        panel.adjustments[0] = 0.06
    except Exception:
        pass
    add_textbox(
        slide, Inches(0.98), Inches(1.7), Inches(11.4), Inches(0.8),
        ["Nothing here is charging yet. Search stays free. Checkout is not live."],
        font=FONT_SANS, size=18, color=FORE, bold=True,
    )
    card(
        slide, Inches(0.72), Inches(2.9), Inches(5.9), Inches(3.55),
        "Chess clubs",
        "Direction of travel: clubs pay to host tournaments and to run teams in the workspace. There is a billing layout in the app. It is not collecting money.",
    )
    card(
        slide, Inches(6.85), Inches(2.9), Inches(5.75), Inches(3.55),
        "Promoted listings",
        "Direction of travel: organizers pay to promote a tournament in search. Ordinary listings would still appear without that. This is not built as a live product.",
        bar_color=BLUE, label_color=BLUE_STRONG,
    )
    notes(slide, "District commercial terms are unset. Happy to answer product questions; packaging and pricing are founder-owned.")

    # 11. Close
    slide = new_slide(prs, 2)
    set_box_text(
        slide.placeholders[0],
        ["That’s the current surface area"],
        font=FONT_DISPLAY, size=36, color=FORE, bold=True,
    )
    set_box_text(
        slide.placeholders[1],
        ["Website and iOS binary are real. Coverage and billing still have gaps. Happy to walk any of it in the product."],
        font=FONT_SANS, size=18, color=MUTED,
    )
    notes(slide, "Stop here. Let them ask.")


def main() -> None:
    assets = write_assets()
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    prs.core_properties.title = "Causey product overview"
    prs.core_properties.subject = "CTO brief for partnerships and marketing"
    prs.core_properties.author = "Causey"
    prs.core_properties.comments = "Internal. Not for external distribution."
    patch_theme(prs)
    patch_master_text_styles(prs)
    add_master_chrome(prs, assets)
    style_layouts(prs)
    build(prs, assets)
    prs.save(str(OUT))
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
