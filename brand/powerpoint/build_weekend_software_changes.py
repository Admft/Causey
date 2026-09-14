#!/usr/bin/env python3
"""Short founder recap of software shipped 11–13 Sep 2026. Non-technical. Max 6 slides."""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_template import (  # noqa: E402
    BLUE,
    BLUE_STRONG,
    FONT_DISPLAY,
    FONT_SANS,
    FORE,
    MUTED,
    RED,
    ROOT,
    SLIDE_H,
    SLIDE_W,
    add_bullets,
    add_master_chrome,
    add_textbox,
    apply_slide_footer_fields,
    card,
    patch_master_text_styles,
    patch_theme,
    set_box_text,
    style_layouts,
    write_assets,
)

OUT = ROOT / "Causey weekend software changes.pptx"


def notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def new_slide(prs, layout_idx=6):
    slide = prs.slides.add_slide(prs.slide_layouts[layout_idx])
    apply_slide_footer_fields(slide)
    return slide


def build(prs, assets) -> None:
    # 1. Title
    slide = new_slide(prs, 0)
    slide.shapes.add_picture(
        str(assets["lockup"]), Inches(0.72), Inches(1.2), height=Inches(0.5)
    )
    set_box_text(
        slide.placeholders[0],
        ["What shipped this weekend"],
        font=FONT_DISPLAY,
        size=40,
        color=FORE,
        bold=True,
    )
    set_box_text(
        slide.placeholders[1],
        ["Founders  ·  11–13 September 2026  ·  live on the website  ·  internal"],
        font=FONT_SANS,
        size=16,
        color=MUTED,
    )
    notes(
        slide,
        "Chess is still the working surface. This is not a launch or a finished product.",
    )

    # 2. One idea
    slide = new_slide(prs, 6)
    add_textbox(
        slide,
        Inches(0.72),
        Inches(0.55),
        Inches(11.9),
        Inches(0.35),
        ["THE WEEKEND IN ONE LINE"],
        font=FONT_SANS,
        size=12,
        color=RED,
        bold=True,
    )
    add_textbox(
        slide,
        Inches(0.72),
        Inches(2.15),
        Inches(11.6),
        Inches(2.3),
        [
            "People can finish the jobs we already advertised — claim a school, answer Going, leave a club — without a button that looks successful and then does nothing."
        ],
        font=FONT_DISPLAY,
        size=26,
        color=FORE,
        bold=True,
    )
    add_textbox(
        slide,
        Inches(0.72),
        Inches(4.85),
        Inches(11.6),
        Inches(1.0),
        [
            "Also shipped: district vs school vs coach seats, stronger sign-in, and admin tools so Causey can find and delete workspaces."
        ],
        font=FONT_SANS,
        size=16,
        color=MUTED,
    )
    notes(
        slide,
        "Do not say the product is complete. Say first sessions stop stalling.",
    )

    # 3. District
    slide = new_slide(prs, 1)
    set_box_text(
        slide.placeholders[0],
        ["Schools and districts"],
        font=FONT_DISPLAY,
        size=32,
        color=FORE,
        bold=True,
    )
    add_bullets(
        slide.placeholders[1],
        [
            "The district office, school office, coaches, and helpers now see different work — totals at the district, named students at the school.",
            "After a district is claimed, its schools come with it, so “hand off this school” actually works.",
            "Adding the first school is a form on the main page, not a link that goes nowhere.",
            "Coaches waiting for a group assignment see “wait,” not a loop between overview and roster.",
            "Causey founders can delete a district or a school from admin when a workspace should not exist.",
        ],
        size=17,
    )
    notes(
        slide,
        "Still assisted: Causey sets the district up. No public district signup.",
    )

    # 4. Family + club
    slide = new_slide(prs, 5)
    set_box_text(
        slide.placeholders[0],
        ["Families, students, and clubs"],
        font=FONT_DISPLAY,
        size=28,
        color=FORE,
        bold=True,
    )
    card(
        slide,
        Inches(0.72),
        Inches(1.5),
        Inches(5.9),
        Inches(5.0),
        "Parents and students",
        "Clearing Going works after a coach invite. A family’s own answer does not turn into a fake invite. Join links are for student accounts only. After the event, Family says Attended instead of a blank. Plan only nags about unfinished registration if the student is still Going.",
    )
    card(
        slide,
        Inches(6.85),
        Inches(1.5),
        Inches(5.75),
        Inches(5.0),
        "Clubs and teams",
        "Only club admins can remove people or rotate the join code — coaches no longer see buttons that fail. Staff can leave a club they do not own. Club settings no longer demand a US state. Copy says club or team, not school.",
        bar_color=BLUE,
        label_color=BLUE_STRONG,
    )
    notes(slide, "Not built: pairings, dues, or a public club directory.")

    # 5. Listings / accounts / Causey desk
    slide = new_slide(prs, 1)
    set_box_text(
        slide.placeholders[0],
        ["Events, sign-in, and Causey’s desk"],
        font=FONT_DISPLAY,
        size=32,
        color=FORE,
        bold=True,
    )
    add_bullets(
        slide.placeholders[1],
        [
            "Editing an already-approved public event no longer quietly takes it down. Drafts stay on the draft list. “Today” follows the person’s timezone.",
            "Search by zip and distance works across chess, debate, STEM, arts, and writing — chess is still the dense directory; the others stay thin.",
            "Sign-in, password reset, and claim-link signup are harder for bots and clearer when a link is expired or the email is wrong.",
            "Causey admin can search users by district or school, see who created an organization, and clear junk problem reports without emailing anyone.",
            "Buttons that copy a claim link vs an activation code, or save vs still-need-to-register, now confirm the action that was actually pressed.",
        ],
        size=16,
    )
    notes(slide, "Do not claim other competition types are complete.")

    # 6. Close — still not
    slide = new_slide(prs, 5)
    set_box_text(
        slide.placeholders[0],
        ["What this is not"],
        font=FONT_DISPLAY,
        size=32,
        color=FORE,
        bold=True,
    )
    card(
        slide,
        Inches(0.72),
        Inches(1.5),
        Inches(5.9),
        Inches(5.0),
        "Still unfinished",
        "No self-serve district signup. No public school directory. No in-app payments. No pairings or messaging. Email at real school volume is unproven. Price, contract, and privacy certification are not product.",
    )
    card(
        slide,
        Inches(6.85),
        Inches(1.5),
        Inches(5.75),
        Inches(5.0),
        "What to do next",
        "Walk one district claim, one family Going/Clear, and one club roster on the live site. If those three finish, this weekend stuck. If any stall, it is an ops or data issue, not a missing slide.",
        bar_color=BLUE,
        label_color=BLUE_STRONG,
    )
    notes(
        slide,
        "Ops still needs the three weekend database updates applied on live Supabase. Do not say that in the room unless someone asks why a live path still fails.",
    )


def main() -> None:
    assets = write_assets()
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    prs.core_properties.title = "What shipped this weekend"
    prs.core_properties.subject = "Founder recap 11–13 September 2026"
    prs.core_properties.author = "Causey"
    prs.core_properties.comments = "Internal. Not for external distribution."
    patch_theme(prs)
    patch_master_text_styles(prs)
    add_master_chrome(prs, assets)
    style_layouts(prs)
    build(prs, assets)
    prs.save(str(OUT))
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
